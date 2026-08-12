from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "UP" / "S019_BV1bstNzUEiU_steel-woman-empowerment_2025-08-07_historical-blind.docx"
SCRIPT_PATH = ROOT / "scripts" / "2025-08-07_steel-woman-empowerment.md"
PREDICTION_PATH = "predictions/2026-08-08_steel-woman-empowerment_historical-blind.md"


THEME = {
    "primary": "1F4D78",
    "primary_dark": "162F4D",
    "accent": "D97706",
    "green": "0F766E",
    "rose": "BE185D",
    "text": "1F2937",
    "gray": "6B7280",
    "line": "D9E2EC",
    "row_alt": "F5F7FA",
    "head_text": "FFFFFF",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9E2EC", sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_table_cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])


def set_run_font(run, size=10.5, bold=False, color="1F2937"):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _get_pbdr(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.insert(0, pBdr)
    return pBdr


def set_para_border_left(p, color="1F4D78", sz=28, space=8):
    pBdr = _get_pbdr(p)
    for el in pBdr.findall(qn("w:left")):
        pBdr.remove(el)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    pBdr.append(left)


def set_para_border_bottom(p, color="D9E2EC", sz=6, space=4):
    pBdr = _get_pbdr(p)
    for el in pBdr.findall(qn("w:bottom")):
        pBdr.remove(el)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def add_para(doc, text="", size=10.5, bold=False, color="1F2937", align=None, before=0, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text):
    p = add_para(doc, text, size=12.5, bold=True, color=THEME["primary"], before=14, after=8)
    set_para_border_left(p, color=THEME["primary"], sz=30, space=8)
    set_para_border_bottom(p, color=THEME["line"], sz=6, space=3)
    return p


def add_subheading(doc, text):
    p = add_para(doc, text, size=10.5, bold=True, color="374151", before=8, after=5)
    set_para_border_left(p, color="9DB8D2", sz=16, space=6)
    return p


def add_table(doc, rows, widths, header_fill=None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_table_borders(table, color=THEME["line"], sz=4)
    set_table_cell_margins(table)
    fill = header_fill or THEME["primary"]
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_shading(cell, fill)
                cell_color = THEME["head_text"]
            else:
                set_cell_shading(cell, THEME["row_alt"] if r_idx % 2 == 0 else "FFFFFF")
                cell_color = THEME["text"]
            cell.text = ""
            for line_idx, line in enumerate(str(value).split("\n")):
                p = cell.paragraphs[0] if line_idx == 0 else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                if line == "":
                    continue
                run = p.runs[0] if p.runs and line_idx == 0 else p.add_run(line)
                run.text = line
                set_run_font(run, size=9, bold=(r_idx == 0), color=cell_color)
    add_para(doc, "", after=2)
    return table


def add_callout(doc, title, body, kind=None):
    if kind is None:
        if "观众旅程" in title:
            kind = "green"
        elif "可以直接拍" in title or "立项" in title:
            kind = "rose"
        else:
            kind = "orange"
    palette = {
        "orange": ("FFF7ED", "D97706", "9A3412", "431407"),
        "green": ("F0FDF9", "0F766E", "134E4A", "042F2E"),
        "rose": ("FDF2F8", "BE185D", "9D174D", "500724"),
    }
    bg, edge, title_c, body_c = palette.get(kind, palette["orange"])
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    set_table_borders(table, color=edge, sz=6)
    set_table_cell_margins(table, top=80, bottom=80, left=140, right=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, bg)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "40")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), edge)
    tc_borders.append(left)
    tc_pr.append(tc_borders)
    cell.text = ""
    p1 = cell.paragraphs[0]
    r1 = p1.runs[0] if p1.runs else p1.add_run(title)
    r1.text = title
    set_run_font(r1, size=11, bold=True, color=title_c)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(body)
    set_run_font(r2, size=10, color=body_c)
    add_para(doc, "", after=2)


def add_gray_note(doc, text):
    return add_para(doc, text, size=8.5, color=THEME["gray"], after=4)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(item)
        set_run_font(run, size=10)



def add_report_header(doc, video_title, meta, title="UP 创作者复盘报告"):
    add_para(doc, title, size=20, bold=True, color=THEME["primary"],
             align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "《" + video_title + "》", size=15, bold=True, color="374151",
             align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, meta, size=10.5, color=THEME["gray"],
             align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_para(doc, "", after=4)
    _t = add_para(doc, "", after=2)
    set_para_border_bottom(_t, color=THEME["accent"], sz=14, space=1)
    add_para(doc, "", after=6)




def clean_script(text):
    lines = text.strip().splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    return "\n".join(line.rstrip() for line in lines).strip()
